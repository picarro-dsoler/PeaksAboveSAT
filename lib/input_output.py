import folium
import geopandas as gpd
from shapely import wkt
from openpyxl import load_workbook
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.hyperlink import Hyperlink
import pandas as pd

#Returns a list of folium maps for each boundary. the boundaryname should e on the list
#the list should be a dictionary with the boundaryname as the key and the folium map as the value
def write_folium(a):
    folium_maps = {}
    for boundary_name, boundary_group in a.groupby('BoundaryName'):
        print(f"Processing boundary: {boundary_name}")

        # Get the centroid of the boundary for this group
        boundary_wkt = boundary_group['boundarywkt4326'].iloc[0]
        boundary_geom = wkt.loads(boundary_wkt)
        centroid = boundary_geom.centroid

        # Calculate bounds of the boundary to determine appropriate zoom level
        bounds = boundary_geom.bounds  # (minx, miny, maxx, maxy)
        sw = [bounds[1], bounds[0]]  # southwest corner [lat, lon]
        ne = [bounds[3], bounds[2]]  # northeast corner [lat, lon]
        m = folium.Map(location=[centroid.y, centroid.x])
        m.fit_bounds([sw, ne])
        boundary_geom = wkt.loads(boundary_wkt)
        
        # Add boundary name label at the centroid
        folium.Marker(
            location=[centroid.y, centroid.x],
            popup=boundary_name,
            tooltip=boundary_name,
            icon=folium.DivIcon(
                html=f'<div style="font-size: 12px; color: black; font-weight: bold; text-align: center; background-color: white; border: 1px solid black; padding: 2px; border-radius: 3px;">{boundary_name}</div>',
                icon_size=(len(boundary_name) * 8, 20),
                icon_anchor=(len(boundary_name) * 4, 10)
            )
        ).add_to(m)
        folium.GeoJson(boundary_geom, name='Boundary').add_to(m)
        for peak_id, peak_row in boundary_group.groupby('PeakId'):
            print(f"  Processing peak: {peak_id}")
            # Get the peak geometry and boundary geometry for this peak
            peak_wkt = peak_row['PeakWkt4326'].iloc[0]
            boundary_wkt = peak_row['boundarywkt4326'].iloc[0]
            emission_rate = peak_row['EmissionRate'].iloc[0]
            # Add peak to the map
            peak_geom = wkt.loads(peak_wkt)
            folium.GeoJson(
                peak_geom, 
                name=f'Peak {peak_id}',
                style_function=lambda x: {'color': 'red', 'weight': 3, 'fillOpacity': 0.7},
                tooltip=folium.Tooltip(f'Peak: {peak_row["PeakName"].iloc[0]}<br>Emission Rate: {emission_rate:.2f} Scf/h<br>Disposition: {peak_row["Disposition"].iloc[0]}')
            ).add_to(m)
        folium_maps[boundary_name] = m
    return folium_maps



def read_excel_with_hyperlinks(self, file_name):
    wb = load_workbook(file_name)
    ws = wb.active
    
    # Read data while preserving hyperlinks
    data = []
    headers = []
    
    # Get headers from first row
    for cell in ws[1]:
        headers.append(cell.value)
    
    # Read data rows
    for row in ws.iter_rows(min_row=2, values_only=False):
        row_data = []
        for cell in row:
            if cell.hyperlink:
                # If cell has hyperlink, store the URL instead of display text
                row_data.append(cell.hyperlink.target)
            else:
                row_data.append(cell.value)
        data.append(row_data)
    
    return pd.DataFrame(data, columns=headers)

def auto_adjust_excel_columns(file_name):
    """
    Auto-adjust column widths in an Excel file to fit content
    """
    wb = load_workbook(file_name)
    ws = wb.active
    
    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        
        # Set column width with some padding
        adjusted_width = min(max_length + 2, 50)  # Cap at 50 characters
        ws.column_dimensions[column_letter].width = adjusted_width
    
    wb.save(file_name)


def write_excel_with_hyperlinks(df, file_name):
    wb = Workbook()
    ws = wb.active
    
    # Write the dataframe to the worksheet
    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)
    
    # Find the Hyperlink column index
    hyperlink_col_idx = None
    for idx, col in enumerate(df.columns, 1):
        if col == 'Hyperlink':
            hyperlink_col_idx = idx
            break
    
    # Convert hyperlink formulas to actual hyperlinks
    if hyperlink_col_idx:
        for row in range(2, ws.max_row + 1):  # Start from row 2 (skip header)
            cell = ws.cell(row=row, column=hyperlink_col_idx)
            if cell.value and isinstance(cell.value, str) and cell.value.startswith('=HYPERLINK('):
                # Extract URL and display text from HYPERLINK formula
                formula = cell.value
                # Parse =HYPERLINK("url", "text")
                start_url = formula.find('"') + 1
                end_url = formula.find('"', start_url)
                url = formula[start_url:end_url]
                
                start_text = formula.find('"', end_url + 1) + 1
                end_text = formula.find('"', start_text)
                display_text = formula[start_text:end_text]
                
                # Set hyperlink
                cell.hyperlink = url
                cell.value = display_text
    
    wb.save(file_name)