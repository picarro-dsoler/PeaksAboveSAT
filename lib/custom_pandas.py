import pandas as pd
import geopandas as gpd
from shapely.geometry import Point
from shapely import wkt
from openpyxl import load_workbook

@pd.api.extensions.register_dataframe_accessor("DA3540")
class DA3540Accessor:
    def __init__(self, pandas_obj):
        self._obj = pandas_obj

    def add_easting_northing(self, longitude_column='PeakGpsLongitude', latitude_column='PeakGpsLatitude', id_column='PeakId'):
        #  Convert the peaks dataframe to a GeoDataFrame
        df = self._obj.copy()
        df['geometry'] = df.apply(lambda row: Point(row[longitude_column], row[latitude_column]), axis=1)
        peaks_gdf = gpd.GeoDataFrame(df, geometry='geometry', crs='EPSG:4326')

        #Add Esting and Northing
        # Reproject to UTM
        geo_data  = peaks_gdf.to_crs('EPSG:32630')

        #Add Esting and Northing    
        bng_gdf = geo_data.to_crs(epsg=27700)

        # --- Step 3: Extract the BNG coordinates (Easting/Northing) from the centroid ---
        # Use the .x and .y of the centroid as you correctly did.
        # This operation should now execute without the RuntimeWarning for most rows.
        bng_gdf['Easting'] = bng_gdf.centroid.x
        bng_gdf['Northing'] = bng_gdf.centroid.y

        # --- Step 4: Update the original dataframe with the new coordinates ---
        # Add the Easting and Northing columns to the original dataframe
        self._obj['Easting'] = bng_gdf['Easting'].values
        self._obj['Northing'] = bng_gdf['Northing'].values


    def assign_boundary_to_peaks(self, boundaries,boundary_wkt_column='boundarywkt4326'):
        #Set the geometry of the peaks dataframe to the PeakWkt4326 column
        df = self._obj.copy()
        df['geometry'] = df['PeakWkt4326'].apply(wkt.loads)
        peaks_gdf = gpd.GeoDataFrame(df, geometry='geometry', crs='EPSG:4326')
        # Convert the boundaries dataframe to a GeoDataFrame
        boundaries['geometry'] = boundaries[boundary_wkt_column].apply(wkt.loads)
        boundaries_gdf = gpd.GeoDataFrame(boundaries, geometry='geometry', crs='EPSG:4326')
        
        # Perform spatial join to find which boundary each peak is located in
        peaks_with_boundaries = gpd.sjoin(peaks_gdf, boundaries_gdf, how='left', predicate='within')
        # For peaks that don't fall within any boundary, try to match using SurveyTag
        unmatched_peaks = peaks_with_boundaries[peaks_with_boundaries['boundaryname'].isna()]
        
        if not unmatched_peaks.empty:
            # Create a mapping from SurveyTag to closest boundary name
            for idx, peak in unmatched_peaks.iterrows():
                survey_tag = peak['SurveyTag']
                if pd.notna(survey_tag):
                    # Find the boundary name that best matches the survey tag
                    best_match = None
                    best_score = 0
                    
                    for boundary_name in boundaries_gdf['boundaryname'].dropna():
                        # Simple string matching - you can improve this logic
                        # Check if survey tag contains parts of boundary name or vice versa
                        tag_lower = str(survey_tag).lower()
                        boundary_lower = str(boundary_name).lower()
                        
                        # Calculate a simple similarity score
                        score = 0
                        if tag_lower in boundary_lower or boundary_lower in tag_lower:
                            score = max(len(tag_lower), len(boundary_lower))
                        else:
                            # Check for common substrings
                            common_parts = set(tag_lower.split()) & set(boundary_lower.split())
                            score = len(common_parts)
                        
                        if score > best_score:
                            best_score = score
                            best_match = boundary_name
                    
                    # Update the boundary name for this peak if we found a match
                    if best_match and best_score > 0:
                        peaks_with_boundaries.loc[idx, 'boundaryname'] = best_match
                        # Also set the corresponding boundary WKT
                        matching_boundary = boundaries_gdf[boundaries_gdf['boundaryname'] == best_match]
                        if not matching_boundary.empty:
                            peaks_with_boundaries.loc[idx, 'boundarywkt4326'] = boundaries_gdf[boundaries_gdf['boundaryname'] == best_match]['boundarywkt4326'].iloc[0]
        #Order with the boundary name where the boundaryname starts with B25, then B25R2
        def sort_key(col):
            # Create a sorting key where B25 comes first, then B25R2, then everything else
            return col.apply(lambda x: 0 if isinstance(x, str) and x.startswith('B25 ') and not x.startswith('B25R2 ') 
                           else 1 if isinstance(x, str) and x.startswith('B25R2 ') 
                           else 2 if isinstance(x, str) and x.startswith('B24 ') 
                           else 3 if isinstance(x, str) and x.startswith('B23 ') 
                           else 4)
        
        p = peaks_with_boundaries.sort_values(by=['PeakId','boundaryname'], key=sort_key, ascending=True, inplace=True)

        #Get the first ocurrcence of the id
        p = peaks_with_boundaries.groupby('PeakId').first().reset_index()
        return p

    def add_unique_boundary_name(self, boundary_name_column='boundaryname'):
        #Remove the first characteres before the - but keep the rest
        #print(self._obj)
        self._obj['UniqueBoundaryName'] = self._obj[boundary_name_column].apply(lambda x: ' - '.join(x.split('-')[1:]) if isinstance(x, str) else None)



    def epoch_to_local_time(self, epoch_column='PeakPlumeEpochStart'):
        #Convert the epoch column to UK date and local time
        self._obj['UTCDateTime'] = pd.to_datetime(self._obj[epoch_column], unit='s', origin='unix', utc=True)
        self._obj['UTCDateTime'] = self._obj['UTCDateTime'].dt.tz_convert('Europe/London')
        #Split the PeakPlumeEpochStart into Date and Time
        self._obj['Date'] = self._obj['UTCDateTime'].dt.strftime('%Y-%m-%d')
        self._obj['LocalTime'] = self._obj['UTCDateTime'].dt.strftime('%H:%M:%S')

    def read_excel_with_hyperlinks(self, file_name):
        wb = load_workbook(file_name)
        ws = wb.active
        
        # Read data while preserving hyperlinks
        data = []
        headers = []
        
        # Get headers from first row
        for cell in ws[1]:
            headers.append(cell.value)
        
        # Read data rows
        for row in ws.iter_rows(min_row=2, values_only=False):
            row_data = []
            for cell in row:
                if cell.hyperlink:
                    # If cell has hyperlink, store the URL instead of display text
                    row_data.append(cell.hyperlink.target)
                else:
                    row_data.append(cell.value)
            data.append(row_data)
        
        box_data = pd.DataFrame(data, columns=headers)